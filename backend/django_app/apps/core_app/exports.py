import os
"""
Export and file import utilities for CSV, Excel (.xlsx), and file parsing.
Preserves existing Flask export capabilities.
"""
import csv
import io
from typing import Any, Tuple, Optional, List, Dict
from django.http import HttpResponse


def generate_csv_response(headers: list[str], rows: list[list[Any]], filename: str) -> HttpResponse:
    """Generate a downloadable CSV response with UTF-8 BOM encoding."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([str(item) if item is not None else "" for item in row])

    csv_data = "\ufeff" + output.getvalue()  # Add UTF-8 BOM for Excel compatibility
    response = HttpResponse(csv_data, content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    return response


def generate_excel_response(
    sheet_title: str,
    headers: list[str],
    rows: list[list[Any]],
    filename: str,
) -> HttpResponse:
    """Generate a formatted Excel (.xlsx) file using openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel limits sheet names to 31 chars

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    alignment = Alignment(horizontal="left", vertical="center")

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = alignment

    # Data rows
    for row in rows:
        clean_row = []
        for val in row:
            if val is None:
                clean_row.append("")
            elif isinstance(val, (int, float)):
                clean_row.append(val)
            else:
                clean_row.append(str(val))
        ws.append(clean_row)

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def generate_json_response(data: Any, filename: str) -> HttpResponse:
    """Generate a formatted, downloadable JSON response."""
    import json
    json_bytes = json.dumps(data, indent=2, default=str).encode("utf-8")
    response = HttpResponse(json_bytes, content_type="application/json; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}.json"'
    return response


def parse_uploaded_file(uploaded_file) -> Tuple[Optional[List[str]], Optional[List[Dict[str, str]]], Optional[str]]:
    """
    Parse an uploaded CSV, XLSX, or JSON file from Django request.FILES.
    Validates file extension, size limit (max 10MB), and returns (headers, list_of_dict_rows, error_message).
    """
    import json
    # Security: File size limit (10MB)
    max_size = 10 * 1024 * 1024
    if getattr(uploaded_file, "size", 0) > max_size:
        return None, None, "File size exceeds maximum allowed limit of 10MB."

    raw_filename = getattr(uploaded_file, "name", "").strip()
    # Security: Sanitize filename to prevent path traversal
    filename = os.path.basename(raw_filename).lower()
    allowed_exts = (".csv", ".xlsx", ".xls", ".json")
    if not any(filename.endswith(ext) for ext in allowed_exts):
        return None, None, "Invalid file format. Please upload a .csv, .xlsx, or .json file."

    rows_data = []
    try:
        if filename.endswith(".csv"):
            file_bytes = uploaded_file.read()
            # Try utf-8-sig first (handles BOM), then fallback to utf-8, latin-1
            text = None
            for enc in ["utf-8-sig", "utf-8", "latin-1"]:
                try:
                    text = file_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                return None, None, "Could not decode file. Please ensure it is saved as UTF-8."

            stream = io.StringIO(text, newline=None)
            reader = csv.DictReader(stream)
            headers = [h.strip() for h in (reader.fieldnames or [])]
            for row in reader:
                clean_row = {k.strip(): (v.strip() if v else "") for k, v in row.items() if k}
                rows_data.append(clean_row)

        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            from openpyxl import load_workbook
            file_bytes = uploaded_file.read()
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return None, None, "Uploaded Excel sheet is empty."

            raw_headers = [str(h or "").strip() for h in rows[0]]
            headers = [h for h in raw_headers if h]

            for row_vals in rows[1:]:
                if not any(row_vals):
                    continue
                row_dict = {}
                for idx, val in enumerate(row_vals):
                    if idx < len(headers):
                        row_dict[headers[idx]] = str(val if val is not None else "").strip()
                rows_data.append(row_dict)

        elif filename.endswith(".json"):
            file_bytes = uploaded_file.read()
            text = file_bytes.decode("utf-8")
            raw_json = json.loads(text)
            records = raw_json
            if isinstance(raw_json, dict):
                # If wrapped in a data or records key
                for candidate in ["data", "records", "items", "rows"]:
                    if candidate in raw_json and isinstance(raw_json[candidate], list):
                        records = raw_json[candidate]
                        break
                else:
                    records = [raw_json]
            if not isinstance(records, list) or not records:
                return None, None, "JSON file does not contain a list of records."

            headers_set = []
            for item in records:
                if isinstance(item, dict):
                    for k in item.keys():
                        if k not in headers_set:
                            headers_set.append(k)

            headers = headers_set
            for item in records:
                if isinstance(item, dict):
                    row_dict = {str(k): str(v if v is not None else "").strip() for k, v in item.items()}
                    rows_data.append(row_dict)

        return headers, rows_data, None

    except Exception as e:
        import logging
        logging.getLogger("bizflow.exports").error("File parsing error: %s", e)
        return None, None, "Error parsing uploaded file: The file format is invalid or corrupted."
