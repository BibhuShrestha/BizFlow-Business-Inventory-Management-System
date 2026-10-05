"""
ExportService — Generate CSV and Excel responses from data.
Migrated from Flask's app/utils/excel_csv.py.
"""
import csv
import io
import logging
from typing import Any

logger = logging.getLogger(__name__)


class ExportService:
    """Generates downloadable file responses."""

    @staticmethod
    def to_csv(headers: list[str], rows: list[list[Any]], filename: str) -> bytes:
        """
        Generate CSV content as bytes.

        Args:
            headers: Column header labels.
            rows: List of row data (list of values per row).
            filename: Base filename (without extension).

        Returns:
            CSV content as UTF-8 encoded bytes.
        """
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)
        return output.getvalue().encode("utf-8-sig")

    @staticmethod
    def to_excel(sheet_title: str, headers: list[str], rows: list[list[Any]]) -> bytes:
        """
        Generate a formatted Excel (.xlsx) workbook as bytes.

        Args:
            sheet_title: Name for the Excel sheet.
            headers: Column header labels.
            rows: List of row data.

        Returns:
            Excel workbook content as bytes.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            logger.error("openpyxl not installed. Cannot generate Excel file.")
            raise

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        alignment = Alignment(horizontal="left", vertical="center")

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = alignment

        for row in rows:
            ws.append(row)

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()


class ImportService:
    """Parse uploaded CSV or Excel files into row dictionaries."""

    @staticmethod
    def parse_file(file_bytes: bytes, filename: str) -> tuple[list[str], list[dict], str | None]:
        """
        Parse a CSV or Excel file and return structured data.

        Args:
            file_bytes: Raw file content as bytes.
            filename: Original filename (used to detect extension).

        Returns:
            Tuple of (headers, rows, error_message).
            error_message is None on success.
        """
        filename_lower = filename.lower()

        if filename_lower.endswith(".csv"):
            return ImportService._parse_csv(file_bytes)
        elif filename_lower.endswith((".xlsx", ".xls")):
            return ImportService._parse_excel(file_bytes)
        else:
            return [], [], "Invalid file format. Please upload a .csv or .xlsx file."

    @staticmethod
    def _parse_csv(file_bytes: bytes) -> tuple[list[str], list[dict], str | None]:
        try:
            stream = io.StringIO(file_bytes.decode("utf-8-sig"), newline=None)
            reader = csv.DictReader(stream)
            headers = [h.strip() for h in (reader.fieldnames or [])]
            rows = [{k.strip(): (v.strip() if v else "") for k, v in row.items() if k} for row in reader]
            return headers, rows, None
        except Exception as exc:
            return [], [], f"Error parsing CSV: {exc}"

    @staticmethod
    def _parse_excel(file_bytes: bytes) -> tuple[list[str], list[dict], str | None]:
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                return [], [], "Uploaded Excel sheet is empty."

            raw_headers = [str(h or "").strip() for h in all_rows[0]]
            headers = [h for h in raw_headers if h]

            rows = []
            for row_vals in all_rows[1:]:
                if not any(row_vals):
                    continue
                row_dict = {}
                for idx, val in enumerate(row_vals):
                    if idx < len(headers):
                        row_dict[headers[idx]] = str(val or "").strip()
                rows.append(row_dict)

            return headers, rows, None
        except Exception as exc:
            return [], [], f"Error parsing Excel: {exc}"
